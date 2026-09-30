"""Test de fumée de l'app empaquetée : lance le programme fabriqué par PyInstaller en mode serveur seul
et vérifie qu'il sert l'interface et sait lire des PDF, Word et PowerPoint (bibliothèques bien incluses).

    python packaging/smoke_test.py dist/Pirouette.app/Contents/MacOS/Pirouette
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
import uuid
from pathlib import Path


def minimal_pdf(text: str) -> bytes:
    """Un PDF d'une page contenant du texte, écrit à la main (aucune dépendance)."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def office_files() -> list[tuple[str, bytes]]:
    import docx
    from pptx import Presentation

    doc = docx.Document()
    doc.add_paragraph("La mitochondrie produit l'ATP.")
    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Le cycle de Calvin"
    deck_bytes = io.BytesIO()
    deck.save(deck_bytes)
    return [("cours.docx", doc_bytes.getvalue()), ("cours.pptx", deck_bytes.getvalue())]


def request(url: str, method: str = "GET", data: bytes | None = None, headers: dict | None = None):
    req = urllib.request.Request(url, data=data, method=method, headers=headers or {})
    with urllib.request.urlopen(req, timeout=30) as response:
        return response.status, response.read()


def multipart(files: list[tuple[str, bytes]]) -> tuple[bytes, str]:
    boundary = uuid.uuid4().hex
    body = io.BytesIO()
    for name, content in files:
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"files\"; filename=\"{name}\"\r\n"
                   f"Content-Type: application/octet-stream\r\n\r\n".encode())
        body.write(content + b"\r\n")
    body.write(f"--{boundary}--\r\n".encode())
    return body.getvalue(), f"multipart/form-data; boundary={boundary}"


def main(binary: str) -> None:
    with tempfile.TemporaryDirectory() as data_dir:
        env = os.environ | {"QUIZZ_DATA_DIR": data_dir}
        proc = subprocess.Popen([binary, "--server-only"], stdout=subprocess.PIPE, env=env, text=True)
        try:
            info = json.loads(proc.stdout.readline())
            base = info["url"].rstrip("/")
            assert info["data_dir"] == data_dir, info

            for path in ("/", "/static/app.js", "/static/style.css", "/static/brand/pirouette.svg",
                         "/manifest.webmanifest", "/api/config", "/api/settings"):
                status, _ = request(base + path)
                assert status == 200, (path, status)
            settings = json.loads(request(base + "/api/settings")[1])
            assert settings["desktop"] is True, settings

            status, body = request(base + "/api/courses", "POST", json.dumps({"name": "Test"}).encode(),
                                   {"Content-Type": "application/json"})
            course_id = json.loads(body)["id"]
            files = [("cours.pdf", minimal_pdf("La photosynthese a lieu dans le chloroplaste.")), *office_files()]
            data, content_type = multipart(files)
            status, body = request(f"{base}/api/courses/{course_id}/files", "POST", data, {"Content-Type": content_type})
            result = json.loads(body)
            assert set(result["results"]) == {"cours.pdf", "cours.docx", "cours.pptx"}, result
            chars = {f["name"]: f["chars"] for f in result["course"]["files"]}
            assert all(n > 10 for n in chars.values()), chars
            # Figures pour Claude : une page de PDF rendue en image (pdfium bien inclus dans l'app).
            pdf_id = next(f["id"] for f in result["course"]["files"] if f["name"] == "cours.pdf")
            status, image = request(f"{base}/api/courses/{course_id}/figure?ref={pdf_id}:1")
            assert status == 200 and image[:2] == b"\xff\xd8", (status, image[:8])
            print(f"Page de PDF rendue en image : {len(image)} octets")
            # Un vrai fichier Pages est lu directement, sans l'app Pages.
            pages = Path(__file__).resolve().parent.parent / "tests" / "data" / "iwork" / "testPages2013.pages"
            data, content_type = multipart([("vrai.pages", pages.read_bytes())])
            status, body = request(f"{base}/api/courses/{course_id}/files", "POST", data, {"Content-Type": content_type})
            read = {f["name"]: f["chars"] for f in json.loads(body)["course"]["files"]}
            assert read.get("vrai.pages", 0) > 500, read
            print(f"Fichier Pages lu directement : {read['vrai.pages']} caractères")

            # Pages n'est pas installé sur les Mac de GitHub : le message doit le dire, sans rester bloqué.
            data, content_type = multipart([("cours.pages", b"PK\x03\x04 pas un vrai fichier")])
            try:
                request(f"{base}/api/courses/{course_id}/files", "POST", data, {"Content-Type": content_type})
                raise AssertionError("un .pages illisible aurait dû être refusé")
            except urllib.error.HTTPError as exc:
                detail = json.loads(exc.read())["detail"]
                assert exc.code == 400 and ("installée" in detail or "convertir" in detail), detail
                print(f"Fichier Pages sans l'app Pages : {detail}")

            print(f"OK : interface servie, PDF/Word/PowerPoint lus ({chars}), données dans {data_dir}")
        finally:
            proc.terminate()
            proc.wait(timeout=10)

        # Rappel quotidien : l'app lancée avec --remind (par macOS) fait son travail et s'arrête, sans fenêtre.
        done = subprocess.run([binary, "--remind"], env=env, capture_output=True, text=True, timeout=120)
        assert done.returncode == 0, done.stderr
        print("Rappel quotidien (--remind) : OK")

        # Branché à l'app Claude : l'app lancée avec --mcp répond au protocole MCP, et rien d'autre sur sa sortie.
        messages = [{"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {"protocolVersion": "2025-06-18"}},
                    {"jsonrpc": "2.0", "method": "notifications/initialized"},
                    {"jsonrpc": "2.0", "id": 2, "method": "tools/call", "params": {"name": "pirouette_cours", "arguments": {}}}]
        talk = subprocess.run([binary, "--mcp"], env=env, capture_output=True, text=True, timeout=120,
                              input="".join(json.dumps(m) + "\n" for m in messages))
        lines = talk.stdout.splitlines()
        assert len(lines) == 2, (talk.stdout, talk.stderr)
        assert json.loads(lines[0])["result"]["serverInfo"]["name"] == "pirouette", lines[0]
        assert "Test" in json.loads(lines[1])["result"]["content"][0]["text"], lines[1]
        print("App Claude (--mcp) : OK")


if __name__ == "__main__":
    main(sys.argv[1])
