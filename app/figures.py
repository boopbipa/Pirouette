"""Les figures des cours (schémas, images, graphiques), pour Claude.

Le texte extrait d'un cours perd ses images. Pour ne pas les envoyer toutes (très coûteux), Pirouette les repère et
place dans le texte lu par Claude un repère court « [Figure f1a2b3c4:12] » à l'endroit où elles se trouvent ; Claude
ne demande l'image (outil pirouette_figure) que si elle l'aide vraiment.

- PDF : une page « a une figure » si elle contient une grande image, ou un schéma dessiné (beaucoup de traits).
  La figure est la page entière, rendue en image : les légendes et les flèches restent autour du schéma.
- PowerPoint : les images posées sur la diapo (clé « 12.1 », « 12.2 »…).
Les images répétées sur la plupart des pages (logo, fond du modèle) ne comptent pas.

Le repérage est fait une fois par fichier, puis gardé à côté du fichier (<id>.figures.json).
"""

from __future__ import annotations

import hashlib
import io
import re
from collections import Counter
from pathlib import Path

INDEX_VERSION = 1
MAX_SIDE = 1100            # côté le plus long de l'image envoyée (≈ 1 000 à 1 500 jetons pour Claude)
MIN_IMAGE_SHARE = 0.04     # une image doit couvrir au moins 4 % de la page / diapo
MIN_PATHS = 25             # un schéma dessiné : au moins 25 traits…
MIN_DRAWING_SHARE = 0.12   # … sur au moins 12 % de la page
REPEATED_SHARE = 0.6       # présente sur 60 % des pages (et au moins 3) : logo ou fond, ignorée

MARKER = re.compile(r"^\[(Page|Diapo) (\d+)\]\s*$")


def supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in {".pdf", ".pptx"}


# ---------- Repérage ----------

def build_index(filename: str, data: bytes) -> dict:
    """{"version", "kind": "pdf" | "pptx", "pages": {"12": ["12"]} ou {"3": ["3.1", "3.2"]}}."""
    ext = Path(filename).suffix.lower()
    pages: dict[str, list[str]] = {}
    if ext == ".pdf":
        pages = _pdf_pages(data)
    elif ext == ".pptx":
        pages = _pptx_pages(data)
    return {"version": INDEX_VERSION, "kind": ext.lstrip("."), "pages": pages}


def _repeated(keys_per_page: list[set]) -> set:
    counts = Counter(key for keys in keys_per_page for key in keys)
    limit = max(3, REPEATED_SHARE * len(keys_per_page))
    return {key for key, n in counts.items() if n >= limit}


def _pdf_pages(data: bytes) -> dict[str, list[str]]:
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    pdf = pdfium.PdfDocument(data)
    try:
        pages = []
        for index in range(len(pdf)):
            page = pdf[index]
            width, height = page.get_size()
            area = max(width * height, 1)
            images, paths = [], []
            for obj in page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE, raw.FPDF_PAGEOBJ_PATH]):
                left, bottom, right, top = obj.get_bounds()
                box = (round(left), round(bottom), round(right), round(top))
                (images if obj.type == raw.FPDF_PAGEOBJ_IMAGE else paths).append(box)
            pages.append((area, images, paths))
            page.close()
    finally:
        pdf.close()
    repeated = _repeated([set(images) | set(paths) for _, images, paths in pages])
    found = {}
    for number, (area, images, paths) in enumerate(pages, start=1):
        big_image = any(_area(box) / area >= MIN_IMAGE_SHARE for box in images if box not in repeated)
        drawing = [box for box in paths if box not in repeated]
        is_drawing = len(drawing) >= MIN_PATHS and _area(_union(drawing)) / area >= MIN_DRAWING_SHARE
        if big_image or is_drawing:
            found[str(number)] = [str(number)]
    return found


def _area(box) -> float:
    left, bottom, right, top = box
    return max(right - left, 0) * max(top - bottom, 0)


def _union(boxes):
    return (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes))


def _pptx_pictures(data: bytes):
    """Pour chaque diapo : [(empreinte, image, extension)] des images assez grandes."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(io.BytesIO(data))
    area = max((deck.slide_width or 1) * (deck.slide_height or 1), 1)

    def pictures(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from pictures(shape.shapes)
            elif hasattr(shape, "image"):
                try:
                    image = shape.image
                except (ValueError, AttributeError, KeyError):  # image liée (hors du fichier) ou vide
                    continue
                share = ((shape.width or 0) * (shape.height or 0)) / area
                if share >= MIN_IMAGE_SHARE:
                    yield hashlib.sha1(image.blob).hexdigest(), image.blob, image.ext

    return [list(pictures(slide.shapes)) for slide in deck.slides]


def _pptx_pages(data: bytes) -> dict[str, list[str]]:
    slides = _pptx_pictures(data)
    repeated = _repeated([{digest for digest, _, _ in pics} for pics in slides])
    found = {}
    for number, pics in enumerate(slides, start=1):
        keys = [f"{number}.{n}" for n, (digest, _, _) in enumerate(pics, start=1) if digest not in repeated]
        if keys:
            found[str(number)] = keys
    return found


# ---------- Image envoyée ----------

def render(filename: str, data: bytes, key: str) -> tuple[bytes, str]:
    """L'image d'une figure, réduite (JPEG ou PNG) : (contenu, type MIME). ValueError si la clé n'existe pas."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data)
        try:
            number = int(key)
            if not 1 <= number <= len(pdf):
                raise ValueError(key)
            page = pdf[number - 1]
            width, height = page.get_size()
            image = page.render(scale=min(MAX_SIDE / max(width, height, 1), 3)).to_pil()
            page.close()
        finally:
            pdf.close()
        return _encode(image)
    if ext == ".pptx":
        match = re.fullmatch(r"(\d+)\.(\d+)", key)
        slides = _pptx_pictures(data)
        if not match or not 1 <= int(match.group(1)) <= len(slides):
            raise ValueError(key)
        pics = slides[int(match.group(1)) - 1]
        if not 1 <= int(match.group(2)) <= len(pics):
            raise ValueError(key)
        from PIL import Image

        try:
            image = Image.open(io.BytesIO(pics[int(match.group(2)) - 1][1]))
            image.load()
        except Exception as exc:  # EMF, WMF… : formats que Pillow ne sait pas dessiner
            raise ValueError(f"image illisible ({pics[int(match.group(2)) - 1][2]})") from exc
        return _encode(image)
    raise ValueError(key)


def _encode(image) -> tuple[bytes, str]:
    image.thumbnail((MAX_SIDE, MAX_SIDE))
    out = io.BytesIO()
    if image.mode in {"RGBA", "LA", "P"}:  # transparence : PNG, sur fond blanc pour les schémas
        from PIL import Image

        background = Image.new("RGB", image.size, "white")
        rgba = image.convert("RGBA")
        background.paste(rgba, mask=rgba.split()[-1])
        background.save(out, "PNG", optimize=True)
        return out.getvalue(), "image/png"
    image.convert("RGB").save(out, "JPEG", quality=82, optimize=True)
    return out.getvalue(), "image/jpeg"


# ---------- Texte lu par Claude ----------

def add_markers(text: str, file_id: str, pages: dict[str, list[str]]) -> str:
    """Ajoute « [Figure <fichier>:<clé>] » sous le repère « [Page n] » / « [Diapo n] » de chaque page qui en a.

    Une page sans texte n'a pas de repère : ses figures vont sous la page précédente (dans le même extrait)."""
    if not pages:
        return text
    wanted = sorted(int(n) for n in pages)
    out, last = [], None
    for line in text.split("\n"):
        match = MARKER.match(line)
        if match:
            number = int(match.group(2))
            if last is not None:  # pages sans texte entre la précédente et celle-ci
                out += [_marker(file_id, key) for n in wanted if last < n < number for key in pages[str(n)]]
            out.append(line)
            out += [_marker(file_id, key) for key in pages.get(str(number), [])]
            last = number
            continue
        out.append(line)
    return "\n".join(out)


def _marker(file_id: str, key: str) -> str:
    return f"[Figure {file_id}:{key}]"


PAGE_NUMBER = re.compile(r"^\s*(?:page\s*)?\d{1,4}(?:\s*(?:/|sur)\s*\d{1,4})?\s*$", re.IGNORECASE)


def boilerplate(text: str) -> set[str]:
    """Les lignes qui se répètent sur la plupart des pages d'un fichier (en-têtes, pieds de page, nom de la fac…)."""
    blocks = [set()]
    for line in text.split("\n"):
        if MARKER.match(line):
            blocks.append(set())
        elif line.strip():
            blocks[-1].add(_shape(line))
    blocks = [b for b in blocks if b]
    return _repeated(blocks) if len(blocks) >= 4 else set()


def compact(text: str, repeated: set[str]) -> str:
    """Retire les lignes répétées (voir boilerplate) et les numéros de page seuls sur leur ligne. Le reste du texte
    ne change pas : les citations restent exactes."""
    kept = [line for line in text.split("\n")
            if MARKER.match(line) or not line.strip()
            or (_shape(line) not in repeated and not PAGE_NUMBER.match(line))]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()


def _shape(line: str) -> str:
    """Une ligne sans ses chiffres ni sa casse : « Biologie — p. 3 » et « Biologie — p. 4 » se ressemblent."""
    return re.sub(r"\d+", "#", " ".join(line.lower().split()))
