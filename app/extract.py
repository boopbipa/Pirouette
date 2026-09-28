"""Extraction du texte des cours déposés (PDF, Word, PowerPoint, Pages, Keynote, texte)."""

from __future__ import annotations

import io
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from .iwork import IWorkError, _wrap_runs, extract_iwork_text, extract_legacy_xml_text

# Version de l'extraction : 2 = le gras et l'italique sont gardés (Markdown). Les fichiers lus avec une version
# plus ancienne sont relus depuis l'original à l'ouverture du cours.
EXTRACT_VERSION = 2

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".pages", ".key", ".txt", ".md", ".markdown"}

# Fichiers Apple (format fermé) : on demande à l'app correspondante de les convertir en Word / PowerPoint.
IWORK = {
    ".pages": {"app": "Pages", "id": "com.apple.iWork.Pages", "format": "Microsoft Word", "ext": ".docx"},
    ".key": {"app": "Keynote", "id": "com.apple.iWork.Keynote", "format": "Microsoft PowerPoint", "ext": ".pptx"},
}


class UnsupportedFileError(ValueError):
    pass


class ConversionError(ValueError):
    """Un fichier Pages / Keynote n'a pas pu être converti ; le message explique quoi faire."""


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pptx":
        text = _extract_pptx(data)
    elif ext in IWORK:
        converted_ext, converted = _convert_iwork(ext, data)
        return extract_text(f"converti{converted_ext}", converted)
    elif ext in {".txt", ".md", ".markdown"}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFileError(
            f"Format non supporté : {ext or filename}. "
            f"Formats acceptés : {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return _clean(text)


def _convert_iwork(ext: str, data: bytes) -> tuple[str, bytes]:
    """Lit un fichier Pages / Keynote. Renvoie (extension, contenu) à passer de nouveau à extract_text."""
    # 1. Format actuel (depuis 2013) : lecture directe du texte, sans l'app Apple ni autorisation.
    try:
        text = extract_iwork_text(data)
        if text.strip():
            return ".txt", text.encode("utf-8")
    except IWorkError:
        pass

    # 2. Anciens fichiers iWork ('09) : l'aperçu PDF complet du document s'il existe (le plus fidèle),
    #    sinon le texte XML (qui peut contenir des textes de modèle).
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for name in ("QuickLook/Preview.pdf", "preview.pdf"):
                if name in archive.namelist():
                    return ".pdf", archive.read(name)
    except zipfile.BadZipFile:
        pass
    try:
        text = extract_legacy_xml_text(data)
        if text.strip():
            return ".txt", text.encode("utf-8")
    except IWorkError:
        pass

    # 3. En dernier recours, on demande à l'app Apple d'exporter le document (macOS uniquement).
    info = IWORK[ext]
    manual = (f"Sinon, ouvre le fichier dans {info['app']} → Fichier → Exporter vers → "
              f"{'Word' if ext == '.pages' else 'PowerPoint'} ou PDF, puis dépose le fichier exporté.")
    not_installed = f"Pour lire un fichier {info['app']}, l'app {info['app']} doit être installée sur ce Mac. {manual}"
    if sys.platform != "darwin" or _app_location(info["id"]) is None:
        raise ConversionError(not_installed)

    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / f"cours{ext}"
        target = Path(tmp) / f"cours{info['ext']}"
        source.write_bytes(data)
        script = f"""
            with timeout of 180 seconds
                tell application id "{info['id']}"
                    set leDocument to open (POSIX file "{source}")
                    export leDocument to (POSIX file "{target}") as {info['format']}
                    close leDocument saving no
                end tell
            end timeout"""
        try:
            result = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=200)
        except subprocess.TimeoutExpired as exc:
            raise ConversionError(f"{info['app']} a mis trop de temps à convertir le fichier. {manual}") from exc
        if result.returncode != 0 or not target.exists():
            error = result.stderr.strip()
            if "-1743" in error:  # l'utilisateur n'a pas (encore) autorisé Pirouette à piloter l'app
                raise ConversionError(
                    f"Pirouette n'a pas l'autorisation d'utiliser {info['app']}. Autorise-la dans "
                    f"Réglages Système → Confidentialité et sécurité → Automatisation → Pirouette → {info['app']}, "
                    f"puis redépose le fichier.")
            if "-10814" in error:  # LaunchServices ne trouve pas l'app
                raise ConversionError(not_installed)
            detail = f" (détail : {error[-200:]})" if error else ""
            raise ConversionError(f"{info['app']} n'a pas pu convertir le fichier. {manual}{detail}")
        return info["ext"], target.read_bytes()


def _app_location(bundle_id: str) -> str | bool | None:
    """Où macOS a installé l'app (quel que soit le dossier) ; None si elle est absente.

    Sans PyObjC (lancement de développement avec run.sh), renvoie True : on laisse AppleScript essayer.
    """
    try:
        from AppKit import NSWorkspace
    except ImportError:
        return True
    url = NSWorkspace.sharedWorkspace().URLForApplicationWithBundleIdentifier_(bundle_id)
    return url.path() if url else None


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        content = page.extract_text() or ""
        if content.strip():
            pages.append(f"[Page {i}]\n{content}")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [_docx_heading_prefix(p) + _docx_emphasis(p) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _docx_emphasis(paragraph) -> str:
    """Texte du paragraphe avec le gras et l'italique en Markdown (**gras**, *italique*, ***les deux***)."""
    if _docx_heading_prefix(paragraph):
        return paragraph.text  # un titre Word est déjà repéré par « # »
    font = paragraph.style.font if paragraph.style is not None else None
    base_bold, base_italic = bool(font and font.bold), bool(font and font.italic)
    chars, flags = [], []
    for run in paragraph.runs:
        flag = (base_bold if run.bold is None else bool(run.bold), base_italic if run.italic is None else bool(run.italic))
        chars.append(run.text)
        flags += [flag] * len(run.text)
    text = "".join(chars)
    if text != paragraph.text:  # liens, champs… : on garde le texte brut
        return paragraph.text
    return _wrap_runs(text, flags)


def _docx_heading_prefix(paragraph) -> str:
    """« # » devant les titres Word (Titre 1, Heading 2…) : garde la structure pour repérer les chapitres."""
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = re.match(r"(?:heading|titre)\s*([1-3])$", style.strip(), re.IGNORECASE)
    return "#" * int(match.group(1)) + " " if match else ""


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Notes : {notes}")
        if texts:
            slides.append(f"[Diapo {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
