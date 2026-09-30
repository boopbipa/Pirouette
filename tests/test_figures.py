import io
import json

from app import figures, mcp_server
from app.chapters import chapter_text
from app.extract import extract_text
from app.storage import Store


def _pdf(pages: list[bytes]) -> bytes:
    """Un PDF écrit à la main : une page par flux de contenu (texte en Helvetica, traits, images en ligne)."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", None, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for stream in pages:
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Contents " + str(len(objects)).encode()
                       + b" 0 R /Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def _text(lines: list[str], number: int) -> bytes:
    body = "".join(f"BT /F1 12 Tf 60 {700 - 20 * i} Td ({line}) Tj ET\n" for i, line in enumerate(lines))
    return (f"BT /F1 9 Tf 60 770 Td (Universite de Lyon - Biologie) Tj ET\n{body}"
            f"BT /F1 9 Tf 290 30 Td ({number}) Tj ET\n").encode("latin-1")


LOGO = b"q 30 0 0 30 540 750 cm BI /W 2 /H 2 /CS /G /BPC 8 ID \x00\xff\xff\x00 EI Q\n"
PHOTO = b"q 400 0 0 300 100 250 cm BI /W 2 /H 2 /CS /G /BPC 8 ID \x10\x80\x80\x10 EI Q\n"
DRAWING = b"".join(f"{100 + 10 * i} 200 m {150 + 10 * i} 600 l S\n".encode() for i in range(30))


def course_pdf() -> bytes:
    return _pdf([
        LOGO + _text(["Chapitre 1 : La cellule", "La cellule est l'unite de base du vivant."], 1),
        LOGO + _text(["La membrane plasmique entoure la cellule."], 2) + PHOTO,
        LOGO + _text(["Chapitre 2 : La mitochondrie", "La mitochondrie produit l'ATP."], 3),
        LOGO + _text(["Le cycle de Krebs a lieu dans la matrice."], 4) + DRAWING,
        LOGO + PHOTO,  # page sans texte
        LOGO + _text(["La chaine respiratoire est dans la membrane interne."], 6),
    ])


def test_pdf_figures_skip_logo_and_find_images_and_drawings():
    index = figures.build_index("cours.pdf", course_pdf())
    assert index["pages"] == {"2": ["2"], "4": ["4"], "5": ["5"]}
    image, mime = figures.render("cours.pdf", course_pdf(), "4")
    assert mime == "image/jpeg" and image[:2] == b"\xff\xd8"
    from PIL import Image
    assert max(Image.open(io.BytesIO(image)).size) == figures.MAX_SIDE


def test_markers_and_compact_text():
    text = extract_text("cours.pdf", course_pdf())
    repeated = figures.boilerplate(text)
    compacted = figures.compact(text, repeated)
    assert "Universite de Lyon" in text and "Universite de Lyon" not in compacted
    assert "\n6\n" not in compacted + "\n" and "La mitochondrie produit" in compacted
    marked = figures.add_markers(compacted, "abcd1234", {"2": ["2"], "4": ["4"], "5": ["5"]})
    lines = marked.split("\n")
    assert lines[lines.index("[Page 2]") + 1] == "[Figure abcd1234:2]"
    # la page 5 (sans texte) arrive juste avant la page 6
    assert lines[lines.index("[Page 6]") - 1] == "[Figure abcd1234:5]"


def test_pptx_pictures_skip_repeated_logo():
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    def png(color, size=(200, 120)):
        out = io.BytesIO()
        Image.new("RGB", size, color).save(out, "PNG")
        out.seek(0)
        return out

    deck = Presentation()
    for i in range(4):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Diapo {i + 1}"
        slide.shapes.add_picture(png("black", (40, 40)), Inches(8.5), Inches(0.2), Inches(1.4), Inches(1.4))  # logo
        if i == 2:
            slide.shapes.add_picture(png("green"), Inches(1), Inches(2), Inches(6), Inches(4))
    data = io.BytesIO()
    deck.save(data)
    index = figures.build_index("cours.pptx", data.getvalue())
    assert index["pages"] == {"3": ["3.1"]}
    image, mime = figures.render("cours.pptx", data.getvalue(), "3.1")
    assert mime == "image/jpeg"


def _call(store, name, **arguments):
    reply = mcp_server.handle(mcp_server.Pirouette(store), {"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                                                            "params": {"name": name, "arguments": arguments}})
    return reply["result"]["content"], reply["result"]["isError"]


def test_claude_reads_markers_and_asks_for_a_figure(tmp_path):
    store = Store(tmp_path / "data")
    course = store.create_course("Bio")
    data = course_pdf()
    text = extract_text("cours.pdf", data)
    lines = text.split("\n")
    chapters = [{"title": "Chapitre 1 : La cellule", "line": lines.index("Chapitre 1 : La cellule")},
                {"title": "Chapitre 2 : La mitochondrie", "line": lines.index("Chapitre 2 : La mitochondrie")}]
    store.add_file(course["id"], "cours.pdf", data, text, chapters, 2)
    fid = store.get_course(course["id"])["files"][0]["id"]
    content, error = _call(store, "pirouette_lire", cours=course["id"], chapitres=[f"{fid}-1"])
    read = content[0]["text"]
    assert not error and f"[Figure {fid}:4]" in read and f"[Figure {fid}:2]" not in read
    assert "Universite de Lyon" not in read
    assert (tmp_path / "data" / "courses" / course["id"] / "files" / f"{fid}.figures.json").exists()
    content, error = _call(store, "pirouette_figure", cours=course["id"], figures=[f"{fid}:4", f"{fid}:3"])
    assert error and "introuvable" in content[0]["text"]
    content, error = _call(store, "pirouette_figure", cours=course["id"], figures=[f"[Figure {fid}:4]"])
    assert not error and [c["type"] for c in content] == ["text", "image"] and content[1]["mimeType"] == "image/jpeg"
    # Une question peut garder sa figure (affichée dans Pirouette) ; ses citations restent vérifiées sur le texte brut
    question = {"type": "reponse_courte", "kind": "cours", "question": "Ou a lieu le cycle de Krebs ?", "choices": [],
                "answer": "Dans la matrice", "explanation": "", "source": "Le cycle de Krebs a lieu dans la matrice.",
                "figure": f"{fid}:4"}
    content, error = _call(store, "pirouette_creer_quiz", cours=course["id"], titre="Mito", chapitres=[f"{fid}-1"],
                           questions=[question])
    assert not error, content
    saved = store.get_quiz(store.list_quizzes(course["id"])[0]["id"])
    assert saved["questions"][0]["figure"] == f"{fid}:4"
    assert json.dumps(chapter_text(text, chapters, 1))  # le texte enregistré n'a pas changé
