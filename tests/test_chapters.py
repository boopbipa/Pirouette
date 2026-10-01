import io
import json

import pytest
from fastapi.testclient import TestClient

from tests.helpers import split_all
from app import main
from app.chapters import ai_candidates, chapter_text, chapters_from_ai, detect_chapters
from app.extract import extract_text
from app.providers import ollama_provider
from app.storage import Store

FILLER = "La cellule est l'unité de base du vivant. " * 20

COURSE = f"""Biologie cellulaire
Sommaire
Chapitre 1 : La cellule ........ 2
Chapitre 2 : La mitochondrie ........ 5
Chapitre 1 : La cellule
{FILLER}
1. Observer au microscope
2. Colorer
Chapitre 2 : La mitochondrie
{FILLER}
Chapitre 3 : La photosynthèse
{FILLER}"""


def test_keyword_chapters_skip_table_of_contents():
    chapters = detect_chapters(COURSE)
    assert [c["title"] for c in chapters] == ["Chapitre 1 : La cellule", "Chapitre 2 : La mitochondrie",
                                              "Chapitre 3 : La photosynthèse"]
    assert chapters[0]["line"] == 0  # le titre du cours et le sommaire (courts) rejoignent le 1er chapitre
    assert "Colorer" in chapter_text(COURSE, chapters, 0)
    assert "Colorer" not in chapter_text(COURSE, chapters, 1)
    assert chapter_text(COURSE, chapters, 2).startswith("Chapitre 3")


def test_long_introduction_becomes_its_own_part():
    text = f"Introduction\n{FILLER}\nI. Les glucides\n{FILLER}\nII. Les lipides\n{FILLER}"
    assert [c["title"] for c in detect_chapters(text)] == ["Début du document", "I. Les glucides", "II. Les lipides"]


def test_word_headings_are_kept_as_markdown():
    import docx

    document = docx.Document()
    for title in ("La cellule", "La mitochondrie"):
        document.add_heading(title, level=1)
        document.add_paragraph(FILLER)
        document.add_heading("Sous-partie", level=2)
        document.add_paragraph("Détail.")
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("cours.docx", buffer.getvalue())
    assert "# La cellule" in text and "## Sous-partie" in text
    assert [c["title"] for c in detect_chapters(text)] == ["La cellule", "La mitochondrie"]


def test_no_obvious_chapters():
    assert detect_chapters(FILLER) == []
    # Listes numérotées qui recommencent à 1 : ce ne sont pas des chapitres
    assert detect_chapters(f"1. Farine\n2. Oeufs\n{FILLER}\n1. Mélanger\n2. Cuire") == []


def test_repeated_page_headers_are_not_candidates():
    pages = "\n".join(f"[Page {i}]\nBiologie L1 — 2026\nTexte de la page {i} qui parle de biologie." for i in range(1, 6))
    assert all("Biologie L1" not in line for _, line in ai_candidates(pages))


def test_ai_answer_is_validated():
    candidates = ai_candidates(COURSE)
    lines = {line: index for index, line in candidates}
    answer = {"chapters": [
        {"line": lines["Chapitre 2 : La mitochondrie"], "title": "Chapitre 2 — Mitochondrie"},
        {"line": lines["Chapitre 1 : La cellule"], "title": "Chapitre 1 — Cellule"},
        {"line": 99999, "title": "inventé"},
    ]}
    assert [c["title"] for c in chapters_from_ai(COURSE, answer, candidates)] == \
        ["Chapitre 1 — Cellule", "Chapitre 2 — Mitochondrie"]
    assert chapters_from_ai(COURSE, {"chapters": []}, candidates) == []
    assert chapters_from_ai(COURSE, {"oops": 1}, candidates) == []


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "store", Store(tmp_path))
    return TestClient(main.app)


def test_generation_uses_only_chosen_chapters(client, monkeypatch):
    seen = []

    async def fake_cards(course_text, n_cards, language, model=None, on_progress=None, avoid=(), **kwargs):
        seen.append(course_text)
        return [{"cards": [{"front": "Q ?", "back": "R.", "source": course_text}]}]

    monkeypatch.setattr(ollama_provider, "generate_cards", fake_cards)
    cid = client.post("/api/courses", json={"name": "Bio"}).json()["id"]
    client.post(f"/api/courses/{cid}/files", files=[("files", ("bio.txt", COURSE.encode(), "text/plain")),
                                                    ("files", ("autre.txt", b"Annexe sans chapitre.", "text/plain"))])
    split_all(client, cid)
    files = client.get(f"/api/courses/{cid}").json()["files"]
    bio, other = files
    assert len(bio["chapters"]) == 3 and bio["chapters_by"] == "auto" and other["chapters"] == []

    def cards(chapters):
        response = client.post(f"/api/courses/{cid}/cards",
                               data={"provider": "local", "model": "m", "count": "1", "chapters": chapters})
        return response

    result = [json.loads(line) for line in cards(f"{bio['id']}-2").text.splitlines()][-1]["result"]
    assert "photosynthèse" in seen[-1] and "mitochondrie" not in seen[-1] and "Annexe" not in seen[-1]
    assert result["cards"][0]["scope"] == ["Chapitre 3 : La photosynthèse"]

    cards(f"{bio['id']}-0,{other['id']}")
    assert "Colorer" in seen[-1] and "Annexe" in seen[-1] and "photosynthèse" not in seen[-1]

    cards("")  # rien de précisé : tout le cours
    assert "photosynthèse" in seen[-1] and "Annexe" in seen[-1]

    assert cards("inconnu-1").status_code == 400  # rien de reconnu : refus clair


def test_ai_detection_endpoint(client, monkeypatch):
    async def fake_ask(system, schema, prompt, model=None):
        numbered = dict(line.split(": ", 1) for line in prompt.split("<titres>\n")[1].split("\n</titres>")[0].split("\n"))
        pick = [int(n) for n, text in numbered.items() if text.startswith("Chapitre") and "...." not in text]
        return {"chapters": [{"line": n, "title": f"Partie {i}"} for i, n in enumerate(pick[-3:], start=1)]}

    monkeypatch.setattr(ollama_provider, "ask", fake_ask)
    cid = client.post("/api/courses", json={"name": "Bio"}).json()["id"]
    client.post(f"/api/courses/{cid}/files", files=[("files", ("bio.txt", COURSE.encode(), "text/plain"))])
    fid = client.get(f"/api/courses/{cid}").json()["files"][0]["id"]
    body = client.post(f"/api/courses/{cid}/files/{fid}/chapters", data={"provider": "local", "model": "m"}).json()
    assert body["found"] == 3
    entry = body["course"]["files"][0]
    assert entry["chapters_by"] == "ai" and [c["title"] for c in entry["chapters"]] == ["Partie 1", "Partie 2", "Partie 3"]


def test_old_files_get_chapters_when_course_opens(client):
    cid = client.post("/api/courses", json={"name": "Bio"}).json()["id"]
    client.post(f"/api/courses/{cid}/files", files=[("files", ("bio.txt", COURSE.encode(), "text/plain"))])
    course = main.store.get_course(cid)
    for f in course["files"]:
        del f["chapters"], f["chapters_by"]
    main.store._save_course(course)
    assert len(client.get(f"/api/courses/{cid}").json()["files"][0]["chapters"]) == 3
