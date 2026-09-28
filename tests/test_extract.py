import io

import docx
import pytest
from pptx import Presentation

from app.extract import UnsupportedFileError, extract_text


def test_txt():
    assert extract_text("cours.md", "# Titre\n\n\n\nContenu".encode()) == "# Titre\n\nContenu"


def test_docx():
    document = docx.Document()
    document.add_paragraph("La mitochondrie produit l'ATP.")
    buffer = io.BytesIO()
    document.save(buffer)
    assert "mitochondrie" in extract_text("cours.docx", buffer.getvalue())


def test_pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "La photosynthèse"
    buffer = io.BytesIO()
    presentation.save(buffer)
    text = extract_text("cours.pptx", buffer.getvalue())
    assert "[Diapo 1]" in text and "photosynthèse" in text


def test_unsupported():
    with pytest.raises(UnsupportedFileError):
        extract_text("image.png", b"...")


# ---------- Fichiers Pages / Keynote ----------
import re
import subprocess
import zipfile
from pathlib import Path

from app import extract
from app.extract import ConversionError


def _docx_bytes(text):
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pages_with_embedded_pdf_preview():
    import importlib.util

    spec = importlib.util.spec_from_file_location("smoke", Path(__file__).parent.parent / "packaging" / "smoke_test.py")
    smoke = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(smoke)  # réutilise le PDF d'une page écrit à la main
    minimal_pdf = smoke.minimal_pdf

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("QuickLook/Preview.pdf", minimal_pdf("Le neurone transmet l'influx nerveux."))
        z.writestr("index.xml", "<doc/>")
    assert "neurone" in extract_text("cours.pages", archive.getvalue())


@pytest.fixture(autouse=True)
def pages_found(monkeypatch):
    monkeypatch.setattr(extract, "_app_location", lambda bundle_id: "/Applications/Pages.app")


def test_pages_converted_by_the_pages_app(monkeypatch):
    calls = []

    def fake_osascript(cmd, **kwargs):
        script = cmd[-1]
        calls.append(script)
        target = re.search(r'export leDocument to \(POSIX file "([^"]+)"\)', script).group(1)
        Path(target).write_bytes(_docx_bytes("La synapse relie deux neurones."))
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(extract.sys, "platform", "darwin")
    monkeypatch.setattr(extract.subprocess, "run", fake_osascript)
    text = extract_text("Neurosciences.pages", b"fichier iwork moderne (pas un zip lisible)")
    assert "synapse" in text
    assert 'tell application id "com.apple.iWork.Pages"' in calls[0] and "as Microsoft Word" in calls[0]


def test_keynote_converted_to_powerpoint(monkeypatch):
    def fake_osascript(cmd, **kwargs):
        target = re.search(r'export leDocument to \(POSIX file "([^"]+)"\)', cmd[-1]).group(1)
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Le potentiel d'action"
        deck.save(target)
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(extract.sys, "platform", "darwin")
    monkeypatch.setattr(extract.subprocess, "run", fake_osascript)
    assert "potentiel d'action" in extract_text("cours.key", b"...")


def test_pages_permission_refused(monkeypatch):
    monkeypatch.setattr(extract.sys, "platform", "darwin")
    monkeypatch.setattr(extract.subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(
        cmd, 1, "", "execution error: Not authorized to send Apple events to Pages. (-1743)"))
    with pytest.raises(ConversionError, match="Automatisation"):
        extract_text("cours.pages", b"...")


def test_pages_without_the_app(monkeypatch):
    monkeypatch.setattr(extract.sys, "platform", "darwin")
    monkeypatch.setattr(extract.subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(
        cmd, 1, "", "LSOpenURLsWithRole() failed with error -10814"))
    with pytest.raises(ConversionError, match="doit être installée"):
        extract_text("cours.pages", b"...")


def test_pages_unknown_error_shows_detail(monkeypatch):
    monkeypatch.setattr(extract.sys, "platform", "darwin")
    monkeypatch.setattr(extract.subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(
        cmd, 1, "", "execution error: Pages got an error: The document couldn’t be exported. (-10000)"))
    with pytest.raises(ConversionError, match="détail : .*-10000"):
        extract_text("cours.pages", b"...")


def test_pages_outside_macos():
    with pytest.raises(ConversionError, match="Exporter vers"):
        extract_text("cours.pages", b"...")


def test_pages_not_found_by_macos_skips_applescript(monkeypatch):
    monkeypatch.setattr(extract.sys, "platform", "darwin")
    monkeypatch.setattr(extract, "_app_location", lambda bundle_id: None)
    monkeypatch.setattr(extract.subprocess, "run", lambda *a, **k: pytest.fail("AppleScript ne doit pas être lancé"))
    with pytest.raises(ConversionError, match="doit être installée"):
        extract_text("cours.pages", b"...")


# ---------- Lecture directe des fichiers iWork (sans l'app Apple) ----------
DATA = Path(__file__).parent / "data" / "iwork"


def _no_applescript(monkeypatch):
    monkeypatch.setattr(extract.subprocess, "run", lambda *a, **k: pytest.fail("la lecture directe doit suffire"))


def test_reads_modern_pages_file_directly(monkeypatch):
    _no_applescript(monkeypatch)
    text = extract_text("cours.pages", (DATA / "testPages2013.pages").read_bytes())
    assert "Sample pages document" in text and "Some plain text to parse." in text
    assert "A text box with text." in text


def test_reads_modern_keynote_file_directly(monkeypatch):
    _no_applescript(monkeypatch)
    text = extract_text("cours.key", (DATA / "testKeynote2013.key").read_bytes())
    assert "Some random text for the sake of testability." in text and "A nice note" in text
    assert "Title Text" not in text and "Body Level" not in text  # textes des modèles de diapositives ignorés


def test_reads_2009_files(monkeypatch):
    _no_applescript(monkeypatch)
    assert "Sample pages document" in extract_text("vieux.pages", (DATA / "testPages.pages").read_bytes())
    assert "A sample presentation" in extract_text("vieux.key", (DATA / "testKeynote.key").read_bytes())


def test_reads_package_variant_with_nested_index_zip(monkeypatch):
    _no_applescript(monkeypatch)
    source = zipfile.ZipFile(DATA / "testPages2013.pages")
    inner, outer = io.BytesIO(), io.BytesIO()
    with zipfile.ZipFile(inner, "w") as z:
        for name in source.namelist():
            if name.startswith("Index/"):
                z.writestr(name, source.read(name))
    with zipfile.ZipFile(outer, "w") as z:
        z.writestr("Index.zip", inner.getvalue())
    assert "Sample pages document" in extract_text("paquet.pages", outer.getvalue())


def test_snappy_roundtrip_on_known_vector():
    from app.iwork import snappy_decompress

    # « abcabcabcabc » : 3 octets littéraux puis une copie de 9 octets à distance 3
    compressed = bytes([12, 0x08, ord("a"), ord("b"), ord("c"), 0x22, 0x03, 0x00])  # 0x22 : copie, longueur 9
    assert snappy_decompress(compressed) == b"abcabcabcabc"


def test_pages_bold_and_italic_become_markdown():
    from app.iwork import extract_iwork_text
    from tests.iwork_builder import char_style, pages_file, para_style

    text = ("La neuroscience\nC’est l’étude scientifique du système nerveux.\n"
            "Neuroanatomie\nC’est regarder la structure du cerveau : à quoi il ressemble.")
    body, title, bold = 10, 11, 12
    styles = [(body, 2022, para_style("Corps", bold=False, italic=False)),
              (title, 2022, para_style("Titre défini", bold=True, italic=True, parent=body)),
              (bold, 2021, char_style(bold=True))]
    lines = text.split("\n")
    starts = [sum(len(l) + 1 for l in lines[:i]) for i in range(len(lines))]
    para_entries = [(starts[0], title), (starts[1], body), (starts[2], title), (starts[3], body)]
    at = text.index("à quoi")
    char_entries = [(0, None), (at, bold), (len(text) - 1, None)]
    result = extract_iwork_text(pages_file(text, para_entries, char_entries, styles))
    assert result.split("\n") == [
        "***La neuroscience***", "C’est l’étude scientifique du système nerveux.",
        "***Neuroanatomie***", "C’est regarder la structure du cerveau : **à quoi il ressemble**.",
    ]


def test_word_bold_and_italic_become_markdown():
    import io

    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from app.extract import extract_text

    document = docx.Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("La neuroscience")
    run.bold = run.italic = True
    body = document.add_paragraph("C'est l'étude du système nerveux. On regarde ")
    body.add_run("à quoi il ressemble").bold = True
    body.add_run(".")
    buffer = io.BytesIO()
    document.save(buffer)
    assert extract_text("cours.docx", buffer.getvalue()).split("\n") == [
        "***La neuroscience***", "C'est l'étude du système nerveux. On regarde **à quoi il ressemble**."]
